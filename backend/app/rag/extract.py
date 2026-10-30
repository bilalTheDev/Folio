"""Text extraction with page numbers for PDF, DOCX, and TXT."""

from __future__ import annotations

import io
from dataclasses import dataclass

from docx import Document as DocxDocument
from pypdf import PdfReader

from backend.app.rag.validate import IngestError


@dataclass(frozen=True)
class PageText:
    page: int
    text: str


@dataclass(frozen=True)
class ExtractedDocument:
    pages: list[PageText]
    page_count: int


def _extract_pdf(data: bytes) -> ExtractedDocument:
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise IngestError("unreadable", "Could not read this PDF.") from exc

    pages: list[PageText] = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        pages.append(PageText(page=index, text=text.strip()))

    return ExtractedDocument(pages=pages, page_count=len(pages))


def _extract_docx(data: bytes) -> ExtractedDocument:
    try:
        document = DocxDocument(io.BytesIO(data))
    except Exception as exc:
        raise IngestError("unreadable", "Could not read this DOCX file.") from exc

    # python-docx has no reliable page breaks; estimate pages by paragraph blocks.
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text and p.text.strip()]
    if not paragraphs:
        return ExtractedDocument(pages=[], page_count=0)

    chars_per_page = 3000
    pages: list[PageText] = []
    buffer: list[str] = []
    size = 0
    page_no = 1
    for para in paragraphs:
        buffer.append(para)
        size += len(para)
        if size >= chars_per_page:
            pages.append(PageText(page=page_no, text="\n\n".join(buffer)))
            buffer = []
            size = 0
            page_no += 1
    if buffer:
        pages.append(PageText(page=page_no, text="\n\n".join(buffer)))

    return ExtractedDocument(pages=pages, page_count=len(pages))


def _extract_txt(data: bytes) -> ExtractedDocument:
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            text = None
    else:
        raise IngestError("unreadable", "Could not decode this text file.")

    cleaned = text.strip()
    if not cleaned:
        return ExtractedDocument(pages=[], page_count=0)

    chars_per_page = 3000
    pages: list[PageText] = []
    for index, start in enumerate(range(0, len(cleaned), chars_per_page), start=1):
        pages.append(PageText(page=index, text=cleaned[start : start + chars_per_page]))
    return ExtractedDocument(pages=pages, page_count=len(pages))


def extract_text(file_type: str, data: bytes) -> ExtractedDocument:
    """Extract per-page text for a validated upload."""
    if file_type == "pdf":
        extracted = _extract_pdf(data)
    elif file_type == "docx":
        extracted = _extract_docx(data)
    elif file_type == "txt":
        extracted = _extract_txt(data)
    else:
        raise IngestError("invalid_type", f"Unsupported file type: {file_type}")

    joined = "\n".join(page.text for page in extracted.pages).strip()
    if not joined:
        raise IngestError(
            "no_text",
            "This file has no selectable text. OCR is not supported yet.",
        )
    return extracted
